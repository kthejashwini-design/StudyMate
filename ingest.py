
from pathlib import Path

import chromadb
from pptx import Presentation
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer


# ----------------------------------
# 1. FOLDER AND DATABASE SETTINGS
# ----------------------------------

PPT_FOLDER = Path("presentations")
DB_FOLDER = "./chroma_db"
COLLECTION_NAME = "my_presentations"

# ----------------------------------
# 2. LOAD EMBEDDING MODEL
# ----------------------------------

print("Loading embedding model...")

model = SentenceTransformer("all-MiniLM-L6-v2")

# ----------------------------------
# 3. CONNECT TO CHROMADB
# ----------------------------------

client = chromadb.PersistentClient(
    path=DB_FOLDER
)

collection = client.get_or_create_collection(
    name=COLLECTION_NAME
)

# ----------------------------------
# 4. FUNCTION TO READ PPT FILES
# ----------------------------------

def read_ppt(file_path):

    presentation = Presentation(file_path)

    pages = []

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            # Extract normal text
            if shape.has_text_frame:
                slide_text.append(shape.text)

            # Extract table text
            if shape.has_table:

                for row in shape.table.rows:

                    row_text = [
                        cell.text
                        for cell in row.cells
                    ]

                    slide_text.append(
                        " | ".join(row_text)
                    )

        text = "\n".join(slide_text).strip()

        if text:
            pages.append(
                (slide_number, text)
            )

    return pages


# ----------------------------------
# 5. FUNCTION TO READ PDF FILES
# ----------------------------------

def read_pdf(file_path):

    pdf = PdfReader(file_path)

    pages = []

    for page_number, page in enumerate(
        pdf.pages, start=1
    ):

        text = page.extract_text() or ""

        text = text.strip()

        if text:
            pages.append(
                (page_number, text)
            )

    return pages


# ----------------------------------
# 6. FUNCTION TO SPLIT TEXT
# ----------------------------------

def split_text(text, chunk_size=500, overlap=100):

    chunks = []

    start = 0

    while start < len(text):

        end = start + chunk_size

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start += chunk_size - overlap

    return chunks


# ----------------------------------
# 7. READ ALL PPT AND PDF FILES
# ----------------------------------

all_chunks = []
all_ids = []
all_metadata = []

if not PPT_FOLDER.exists():

    raise FileNotFoundError(
        "Please create the presentations folder "
        "and add your PPT or PDF files."
    )

files = sorted(
    list(PPT_FOLDER.glob("*.pptx"))
    + list(PPT_FOLDER.glob("*.pdf"))
)

if not files:

    raise FileNotFoundError(
        "No PPT or PDF files found in presentations/"
    )


for file_path in files:

    print(f"\nReading: {file_path.name}")

    try:

        # Read PPT or PDF
        if file_path.suffix.lower() == ".pptx":

            pages = read_ppt(file_path)

        else:

            pages = read_pdf(file_path)

        # Process every slide or page
        for page_number, text in pages:

            chunks = split_text(text)

            for chunk_number, chunk in enumerate(
                chunks, start=1
            ):

                chunk_id = (
                    f"{file_path.name}_"
                    f"{page_number}_"
                    f"{chunk_number}"
                )

                all_chunks.append(chunk)

                all_ids.append(chunk_id)

                all_metadata.append({
                    "source": file_path.name,
                    "page": page_number,
                    "chunk": chunk_number
                })

        print(f"Extracted {len(pages)} pages/slides")

    except Exception as error:

        print(
            f"Could not read {file_path.name}: {error}"
        )


# ----------------------------------
# 8. GENERATE EMBEDDINGS
# ----------------------------------

if not all_chunks:

    print("No readable text found.")
    raise SystemExit


print("\nGenerating embeddings...")

embeddings = model.encode(
    all_chunks,
    show_progress_bar=True
).tolist()


# ----------------------------------
# 9. STORE DATA IN CHROMADB
# ----------------------------------

print("\nSaving data to ChromaDB...")

collection.upsert(
    ids=all_ids,
    documents=all_chunks,
    embeddings=embeddings,
    metadatas=all_metadata
)


# ----------------------------------
# 10. DISPLAY RESULTS
# ----------------------------------

print("\nIngestion completed!")

print("Files processed:", len(files))

print("Chunks stored:", len(all_chunks))

print("Collection:", collection.name)

print("Database:", DB_FOLDER)